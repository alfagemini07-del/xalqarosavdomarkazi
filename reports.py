"""Paginated reports and numeric XLSX / embedded-font PDF exports."""
import io
from decimal import Decimal
from xml.sax.saxutils import escape

from flask import abort, g, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from receipt_render import font_path

pdfmetrics.registerFont(TTFont("Tarozi", font_path()))
pdfmetrics.registerFont(TTFont("TaroziBold", font_path(True)))

MODES = {
    "hour": ("Soatlik tushum", "TO_CHAR(w.created_at AT TIME ZONE %s, 'YYYY-MM-DD HH24:00')"),
    "day": ("Kunlik tushum", "TO_CHAR(w.business_date, 'YYYY-MM-DD')"),
    "week": ("Haftalik tushum", "TO_CHAR(DATE_TRUNC('week',w.business_date), 'YYYY-MM-DD')"),
    "month": ("Oylik tushum", "TO_CHAR(w.business_date, 'YYYY-MM')"),
    "operator": ("Operatorlar hisoboti", "COALESCE(u.username, 'Import')"),
    "method": ("To'lov turlari", "CASE w.payment_method WHEN 'cash' THEN 'Naqd' WHEN 'card' THEN 'Karta' ELSE 'Hisob raqam' END"),
    "vehicle": ("Eng ko'p kelgan avtomobillar", "w.plate_number"),
}


def register_reports(app, db, roles_required, now_local, parse_date, timezone_name, company):
    def filters():
        today = now_local().date()
        try:
            start = parse_date(request.args.get("start", today.isoformat()), "Boshlanish")
            end = parse_date(request.args.get("end", today.isoformat()), "Tugash")
            operator = int(request.args.get("operator") or 0)
            page = max(1, int(request.args.get("page", "1")))
            per_page = max(10, min(100, int(request.args.get("per_page", "20"))))
        except (ValueError, TypeError):
            abort(400, description="Hisobot sanasi yoki filtri noto'g'ri")
        if start > end or (end - start).days > 366:
            abort(400, description="Hisobot oralig'i 1–367 kun bo'lishi kerak")
        where, params = ["w.business_date BETWEEN %s AND %s"], [start, end]
        if operator:
            where.append("w.created_by=%s")
            params.append(operator)
        method = request.args.get("method", "")
        if method:
            if method not in {"cash", "card", "bank"}:
                abort(400)
            where.append("w.payment_method=%s")
            params.append(method)
        return start, end, operator, method, " AND ".join(where), params, page, per_page

    def query_report(export=False, export_limit=20000):
        start, end, operator, method, where, params, page, size = filters()
        mode = request.args.get("group", "day")
        base = " FROM weighings w LEFT JOIN users u ON u.id=w.created_by WHERE " + where
        args = list(params)
        if mode in MODES:
            title, expr = MODES[mode]
            if mode == "hour":
                args.insert(0, timezone_name)
            query = ("SELECT " + expr + " AS label,COUNT(*) AS count,COALESCE(SUM(w.price),0) AS total," 
                     "SUM(w.weighing_fee) AS weighing,SUM(w.entry_fee) AS entry,SUM(w.reload_fee) AS reload"
                     + base + " AND w.status='paid' GROUP BY 1")
            headers = ["Guruh", "Avtomobil", "Jami so'm", "Tarozi", "Hudud", "Qayta yuklash"]
            keys = ["label", "count", "total", "weighing", "entry", "reload"]
            order = "count DESC,total DESC,label" if mode == "vehicle" else "total DESC,label" if mode in {"operator", "method"} else "label"
        elif mode == "services":
            title = "Xizmatlar hisoboti"
            query = ("SELECT s.label,COUNT(*) AS count,SUM(s.amount) AS total FROM weighings w "
                     "CROSS JOIN LATERAL (VALUES ('Tarozi',w.weighing_fee,TRUE),"
                     "('Hududga kirish',w.entry_fee,w.entry_service),('Qayta yuklash',w.reload_fee,w.reload_service)) "
                     "s(label,amount,selected) WHERE " + where + " AND w.status='paid' AND s.selected GROUP BY s.label")
            headers, keys, order = ["Xizmat", "Soni", "Jami so'm"], ["label", "count", "total"], "label"
        elif mode == "cancelled":
            title = "Bekor qilingan cheklar"
            query = ("SELECT w.receipt_no,w.plate_number,COALESCE(u.username,'Import') AS operator,"
                     "TO_CHAR(w.cancelled_at AT TIME ZONE %s,'YYYY-MM-DD HH24:MI') AS time,w.price AS amount"
                     + base + " AND w.status='cancelled'")
            args.insert(0, timezone_name)
            headers, keys, order = ["Chek", "Avtomobil", "Operator", "Bekor qilingan vaqt", "Operatsiya summasi"], ["receipt_no", "plate_number", "operator", "time", "amount"], "time DESC,receipt_no"
        elif mode == "reprints":
            title = "Qayta chop urinishlari"
            query = ("SELECT w.receipt_no,w.plate_number,COUNT(p.id) AS attempts,"
                     "COUNT(*) FILTER(WHERE p.state='confirmed') AS confirmed,"
                     "TO_CHAR(MAX(p.created_at) AT TIME ZONE %s,'YYYY-MM-DD HH24:MI') AS time "
                     "FROM weighings w JOIN print_attempts p ON p.weighing_id=w.id WHERE " + where +
                     " GROUP BY w.id HAVING COUNT(p.id)>1")
            args.insert(0, timezone_name)
            headers, keys, order = ["Chek", "Avtomobil", "Jami urinish", "Tasdiqlangan", "Oxirgi urinish"], ["receipt_no", "plate_number", "attempts", "confirmed", "time"], "time DESC,receipt_no"
        else:
            abort(400)
        with db().connection() as conn:
            conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
            total = conn.execute("SELECT COUNT(*) AS n FROM (" + query + ") q", args).fetchone()["n"]
            if export and total > export_limit:
                limit_label = f"{export_limit:,}".replace(",", " ")
                abort(400, description=f"Eksport {limit_label} qatordan oshdi. Sana yoki operator filtrini toraytiring.")
            page = min(page, max(1, (total + size - 1) // size))
            rows = conn.execute(query + " ORDER BY " + order + " LIMIT %s OFFSET %s", args + [max(total, 1) if export else size, 0 if export else (page-1)*size]).fetchall()
            totals = conn.execute(
                "SELECT COUNT(*) FILTER(WHERE w.status='paid') AS paid,"
                "COUNT(*) FILTER(WHERE w.status='cancelled') AS cancelled,"
                "COALESCE(SUM(w.price) FILTER(WHERE w.status='paid'),0) AS total,"
                "COALESCE(SUM(w.price) FILTER(WHERE w.status='paid' AND w.payment_method='cash'),0) AS cash,"
                "COALESCE(SUM(w.price) FILTER(WHERE w.status='paid' AND w.payment_method='card'),0) AS card,"
                "COALESCE(SUM(w.price) FILTER(WHERE w.status='paid' AND w.payment_method='bank'),0) AS bank" + base, params
            ).fetchone()
            operator_row = conn.execute("SELECT username FROM users WHERE id=%s", (operator,)).fetchone() if operator else None
        rows = [[int(r[k]) if isinstance(r[k], Decimal) else r[k] for k in keys] for r in rows]
        return dict(title=title, headers=headers, rows=rows, total=int(total), page=page, per_page=size,
                    pages=max(1,(total+size-1)//size), start=start.isoformat(), end=end.isoformat(),
                    operator=operator, operator_name=operator_row["username"] if operator_row else "Barchasi", method=method, totals={k:int(v) for k,v in totals.items()})

    def pdf_report(title, subtitle, headers, rows, summary="", closing=False):
        output = io.BytesIO()
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "Tarozi"
        styles["Heading1"].fontName = "TaroziBold"
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=30)
        p = lambda value: Paragraph(escape(str(value if value is not None else "—")), styles["Normal"])
        story = [Paragraph(escape(title),styles["Heading1"]),p(company),p(subtitle),p(summary),Spacer(1,14)]
        table = Table([[p(v) for v in headers]] + [[p(v) for v in row] for row in rows], repeatRows=1, colWidths=[doc.width/len(headers)]*len(headers))
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#dbeafe")),("GRID",(0,0),(-1,-1),.4,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
        story.append(table)
        if closing:
            story += [Spacer(1,25),p("Topshirdi: __________________    Qabul qildi: __________________"),p("Imzo: __________________      Imzo: __________________")]
        def footer(canvas, document):
            canvas.setFont("Tarozi",9)
            canvas.drawRightString(815,15,f"Sahifa {document.page}")
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
        output.seek(0)
        return output

    @app.get("/api/admin/analytics")
    @roles_required("admin", "techadmin")
    def analytics():
        return jsonify(success=True, **query_report())

    @app.get("/api/admin/report-operators")
    @roles_required("admin", "techadmin")
    def report_operators():
        with db().connection() as conn:
            rows = conn.execute("SELECT id,username FROM users ORDER BY username").fetchall()
        return jsonify(success=True, users=rows)

    @app.get("/api/admin/analytics/export/<kind>")
    @roles_required("admin", "techadmin")
    def export_analytics(kind):
        if kind not in {"xlsx", "pdf"}:
            abort(404)
        # PDF layout keeps Paragraph objects in memory; bound it separately on small servers.
        data = query_report(export=True, export_limit=2000 if kind == "pdf" else 20000)
        subtitle = f"{data['start']} — {data['end']}; operator: {data['operator_name']}; to'lov: {data['method'] or 'Barchasi'}"
        summary = f"To'langan: {data['totals']['paid']}; tushum: {data['totals']['total']:,} so'm; bekor: {data['totals']['cancelled']}".replace(","," ")
        if kind == "pdf":
            stream = pdf_report(data["title"],subtitle,data["headers"],data["rows"],summary)
            mime = "application/pdf"
        else:
            book = Workbook(); sheet = book.active; sheet.title = "Hisobot"
            sheet.append([data["title"]]); sheet.append([subtitle]); sheet.append([summary]); sheet.append(data["headers"])
            for row in data["rows"]:
                sheet.append(row)
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,str):
                        cell.data_type="s"  # User text must never become an Excel formula.
                    elif isinstance(cell.value,(int,float)):
                        cell.number_format='#,##0'
            for cell in sheet[4]:
                cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="DBEAFE")
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width=26
            sheet.freeze_panes="A5"; sheet.auto_filter.ref=f"A4:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}"
            stream=io.BytesIO(); book.save(stream); stream.seek(0)
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return send_file(stream,mimetype=mime,as_attachment=True,download_name=f"hisobot-{data['start']}-{data['end']}.{kind}")

    @app.get("/api/admin/cash-closing.pdf")
    @roles_required("admin", "techadmin")
    def cash_closing():
        start,end,operator,method,where,params,_,_ = filters()
        if start != end or method:
            abort(400,description="Dalolatnoma uchun bitta kun va barcha to'lov turlarini tanlang")
        try:
            actual = int(request.args["actual_cash"])
            if not 0 <= actual <= 10**15: raise ValueError()
        except (KeyError,ValueError):
            abort(400,description="Kassadagi haqiqiy naqd summani kiriting")
        with db().connection() as conn:
            rows=conn.execute("SELECT payment_method,COUNT(*) AS count,SUM(price) AS total FROM weighings w WHERE "+where+" AND status='paid' GROUP BY payment_method",params).fetchall()
            operator_name = conn.execute("SELECT username FROM users WHERE id=%s",(operator,)).fetchone() if operator else None
        values={r["payment_method"]:int(r["total"]) for r in rows}
        table=[[{"cash":"Naqd","card":"Karta","bank":"Hisob raqam"}[r["payment_method"]],r["count"],int(r["total"])] for r in rows]
        table += [["Jami",sum(r["count"] for r in rows),sum(values.values())],["Haqiqiy naqd", "",actual],["Farq (haqiqiy − hisob)","",actual-values.get("cash",0)]]
        subtitle=f"Sana: {start}; operator: {operator_name['username'] if operator_name else 'Barchasi'}; tuzdi: {g.user['username']}; {now_local():%d.%m.%Y %H:%M:%S}"
        stream=pdf_report("Kunlik kassa yopilish dalolatnomasi",subtitle,["To'lov turi","Soni","So'm"],table,"Yaratilgan paytdagi ma'lumotlar. Keyingi operatsiyalar bu nusxaga kirmaydi.",True)
        return send_file(stream,mimetype="application/pdf",as_attachment=True,download_name=f"kassa-{start}.pdf")
